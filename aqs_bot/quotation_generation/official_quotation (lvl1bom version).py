from docx import Document
from docx.shared import Inches
import pyodbc
import json
import configparser

config = configparser.ConfigParser()
config.read('sql_connect.cfg')

conn = pyodbc.connect('Driver=' + config['database']['driver'] + ';'
                      'Server=' + config['database']['server'] + ';'
                      'Database=' + config['database']['database'] + ';'
                      'Trusted_Connection=' + config['database']['trusted_connection'] + ';')
cursor = conn.cursor()


quotation = cursor.execute(
    "select labour_cost, markup_pct, labour_cost_description from dbo.quotation where quotation_no=?", 'quotation_one')
res = quotation.fetchone()

markup_pct = None
if res.markup_pct != None:
    markup_pct = res.markup_pct/100


row_level = cursor.execute(
    "select id, row, component_no, description, quantity, uom, unit_price, remark, crawl_info, bom_id from dbo.quotation_component where quotation_no=? and bom_id IS NULL ORDER BY row ASC", 'quotation_one')
result = row_level.fetchall()

bom_list = []
current_bom = None
for each_row in result:
    bom_list.append({'bom_id': each_row.id, 'bom_no': each_row.component_no,'row':each_row.row, 'qty':each_row.quantity, 'uom':each_row.uom})

all_bom = []
curr_ind = 0
for bom in bom_list:
    bom_items = []
    bom_items.append(bom['bom_id'])
    row_level = None
    curr_bom_row = bom['row']
    total_price = 0
    try: 
        next_bom_row = bom_list[curr_ind+1]['row']
        row_level = cursor.execute(
        "select id, row, component_no, description, quantity, uom, unit_price, remark, crawl_info, bom_id, is_bom from dbo.quotation_component where quotation_no=? and (row >= ? and row < ?) ORDER BY row ASC", 'quotation_one', curr_bom_row, next_bom_row)
    except:
         row_level = cursor.execute(
        "select id, row, component_no, description, quantity, uom, unit_price, remark, crawl_info, bom_id, is_bom from dbo.quotation_component where quotation_no=? and row >= ? ORDER BY row ASC", 'quotation_one', curr_bom_row)
       
    result = row_level.fetchall()
    for res in result:
        # print(res.row)
        
        if res.bom_id in bom_items or (res.bom_id == None and res.is_bom == 0): #if is a loose item  not belonging to any set.. OR it's a item belonging to a set
            multiple_supplier = False

            bom_items.append(res.id)

            if res.crawl_info != None:
                crawl_info = json.loads(res.crawl_info)
                if len(crawl_info) > 1:
                    for ea_supplier in crawl_info:
                        total_price+= ((ea_supplier['unit_price'] * markup_pct) + ea_supplier['unit_price']) * ea_supplier['qty']
                    multiple_supplier = True
            
            if multiple_supplier == False:
                total_price+= ((res.unit_price * markup_pct) + res.unit_price) * res.quantity

    all_bom.append({'part_no': bom['bom_no'], 'qty': bom['qty'], 'uom': bom['uom'], 'unit_price': total_price/bom['qty'], 'total_price': total_price})
    curr_ind+=1


document = Document('aqs_bot/quotation_generation/Official Quotation Template.docx')

# print(document.tables[1].cell(0,1).text)

table = document.tables[2]
sn = 1
for ea_bom in all_bom:
    row = table.add_row().cells#https://github.com/python-openxml/python-docx/issues/205
    row[0].text = str(sn)
    row[1].text = str(ea_bom['part_no'])
    row[3].text = str(ea_bom['qty'])
    row[4].text = str(ea_bom['uom'])
    row[5].text = str(ea_bom['unit_price'])
    row[6].text = str(ea_bom['total_price'])
    sn+=1

document.save('aqs_bot/quotation_generation/baby.docx')